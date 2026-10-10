#!/usr/bin/env python3
"""Build patients/run_tracker.xlsx: which tracked patients have been run, copied back and reviewed.

Roster + manual columns come from patients/run_tracker.tsv. Automatic columns are read from the
local copies of the farm results (~/Documents/hsc_results/<P>, or the row's results_dir):
  colonies        rows of patients/*/<run_unit>/colonies.tsv (samples.tsv of the run when present)
  results local   <P>.somatic.xlsx (or <P>.annotated.csv.gz alone = report incomplete)
  somatic / tiers the README sheet of <P>.somatic.xlsx
  reviewed        verdicts in <results>/.review/<P>.reviews.sqlite (insertion_review app)
  commit          PT_COMMIT from the run's run.env when the run recorded it (pipeline.sh freeze_env)

Edits Jeremy makes to the manual columns IN THE XLSX win: the hidden _generated sheet keeps what the
last build wrote, a cell that differs from it was edited by hand, and that value is written back to
the TSV before the rebuild. Everything else comes from the TSV.

  python3 cluster/run_tracker.py [--results ~/Documents/hsc_results] [--out patients/run_tracker.xlsx]
"""
import argparse
import csv
import datetime
import glob
import os
import sqlite3
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MANUAL = ["farm_status", "farm_run_dir", "commit", "results_dir", "review_signed_off", "notes"]
FARM_STATUS = ["not run", "submitted", "running", "report incomplete", "done", "failed"]

# (header, key, width)
COLUMNS = [
    ("PD id", "pd", 10),
    ("Role", "role", 18),
    ("Cohort", "cohort", 34),
    ("Why tracked", "reason", 24),
    ("TP53 on tree", "tp53_on_tree", 30),
    ("Run unit", "run_unit", 24),
    ("Assembly", "assembly", 12),
    ("Colonies", "colonies", 9),
    ("Farm run", "farm_status", 15),
    ("Farm run dir", "farm_run_dir", 40),
    ("PEAR-TREE commit", "commit_shown", 34),
    ("Data local", "data_local", 11),
    ("Results dir", "results_shown", 36),
    ("Somatic rows", "somatic", 9),
    ("Tier B / C / D", "tiers", 13),
    ("Excluded", "excluded", 9),
    ("Reviewed", "reviewed", 9),
    ("Review progress", "review_progress", 15),
    ("Review signed off", "review_signed_off", 10),
    ("Notes", "notes", 60),
]
HEADER_KEY = {h: k for h, k, _ in COLUMNS}
KEY_HEADER = {k: h for h, k, _ in COLUMNS}
# manual key -> column key shown in the sheet (commit and results_dir are shown merged with auto values)
SHOWN = {"commit": "commit_shown", "results_dir": "results_shown"}


def read_tsv(path):
    with open(path) as fh:
        lines = [l for l in fh if not l.startswith("#")]
    comments = [l for l in open(path) if l.startswith("#")]
    return comments, list(csv.DictReader(lines, delimiter="\t"))


def write_tsv(path, comments, rows):
    fields = list(rows[0].keys())
    with open(path, "w") as fh:
        fh.writelines(comments)
        w = csv.DictWriter(fh, fields, delimiter="\t", lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def merge_back(xlsx, rows):
    """Copy hand edits of manual columns from the xlsx into rows; returns the edited cells."""
    if not os.path.exists(xlsx):
        return []
    wb = openpyxl.load_workbook(xlsx)
    if "tracker" not in wb.sheetnames or "_generated" not in wb.sheetnames:
        return []

    def table(ws):
        it = ws.iter_rows(values_only=True)
        hdr = next(it)
        return {r[0]: dict(zip(hdr, r)) for r in it if r and r[0]}

    now, gen = table(wb["tracker"]), table(wb["_generated"])
    edits = []
    for row in rows:
        cur, old = now.get(row["pd"]), gen.get(row["pd"])
        if not cur or not old:
            continue
        for key in MANUAL:
            h = KEY_HEADER[SHOWN.get(key, key)]
            a, b = cur.get(h), old.get(h)
            a = "" if a is None else str(a)
            b = "" if b is None else str(b)
            if a != b:
                row[key] = a
                edits.append((row["pd"], key, b, a))
    return edits


def patient_dir(run_unit):
    if not run_unit:
        return None
    hits = glob.glob(os.path.join(REPO, "patients", "*", run_unit))
    return hits[0] if hits else None


def count_colonies(run_unit, pd, results, root):
    for d in filter(None, [results, os.path.join(root, run_unit or pd)]):
        if os.path.exists(os.path.join(d, "samples.tsv")):
            with open(os.path.join(d, "samples.tsv")) as fh:
                return sum(1 for l in fh if l.strip())
    d = patient_dir(run_unit)
    if not d or not os.path.exists(os.path.join(d, "colonies.tsv")):
        return ""
    with open(os.path.join(d, "colonies.tsv")) as fh:
        samples = [l.split("\t")[6].strip() for l in fh
                   if not l.startswith("#") and not l.startswith("donor\t") and l.count("\t") >= 6]
    if not samples:
        trees = glob.glob(os.path.join(d, "*.tree"))
        if trees:
            # placeholder colonies.tsv: the tree's tips (newick leaves = commas + 1)
            return f"{open(trees[0]).read().count(',') + 1} (tree tips)"
        return "pending"
    # pair / two-timepoint folders hold several PD ids: count this PD's colonies
    own = [s for s in samples if s.startswith(pd)]
    return len(own) if own else len(samples)


def find_results(root, row):
    if row["results_dir"]:
        p = os.path.join(root, row["results_dir"])
        return p if os.path.isdir(p) else None
    for name in (row["pd"], row["run_unit"]):
        if name and os.path.isdir(os.path.join(root, name)):
            return os.path.join(root, name)
    return None


def somatic_summary(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    info = {}
    for r in wb["README"].iter_rows(values_only=True):
        if r and len(r) >= 2 and r[0] in ("rows", "tier B", "tier C", "tier D", "excluded"):
            info[r[0]] = int(str(r[1]).split()[0])
    return info


def review_count(results, stem):
    db = os.path.join(results, ".review", f"{stem}.reviews.sqlite")
    if not os.path.exists(db):
        return None
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    try:
        return con.execute("select count(*) from reviews where verdict != ''").fetchone()[0]
    finally:
        con.close()


def run_env_commit(results):
    env = os.path.join(results, "run.env") if results else ""
    if not env or not os.path.exists(env):
        return ""
    for l in open(env):
        if l.startswith("PT_COMMIT="):
            return l.split("=", 1)[1].strip().strip("'")
    return ""


def auto_fields(root, row):
    res = find_results(root, row)
    out = {"colonies": count_colonies(row["run_unit"], row["pd"], res, root),
           "data_local": "no", "results_shown": row["results_dir"], "somatic": "", "tiers": "",
           "excluded": "", "reviewed": "", "review_progress": "", "commit_shown": row["commit"]}
    if not res:
        return out
    out["results_shown"] = os.path.relpath(res, root)
    stem = os.path.basename(res)
    xl = os.path.join(res, f"{stem}.somatic.xlsx")
    if os.path.exists(xl):
        s = somatic_summary(xl)
        out["data_local"] = "yes"
        out["somatic"] = s.get("rows", "")
        out["tiers"] = f"{s.get('tier B', 0)} / {s.get('tier C', 0)} / {s.get('tier D', 0)}"
        out["excluded"] = s.get("excluded", "")
        n = review_count(res, stem)
        if n is None:
            # reviews of an earlier run of the same patient do not carry over: report none
            n = 0
        out["reviewed"] = n
        rows = s.get("rows") or 0
        out["review_progress"] = ("not started" if n == 0 else
                                  "complete" if rows and n >= rows else f"{100 * n // max(rows, 1)} %")
    elif os.path.exists(os.path.join(res, f"{stem}.annotated.csv.gz")):
        out["data_local"] = "partial"
        out["review_progress"] = "no somatic table"
    commit = run_env_commit(res)
    if commit:
        out["commit_shown"] = commit if not row["commit"] else f"{commit} (run.env); {row['commit']}"
    return out


FILL = {
    "done": "C6EFCE", "yes": "C6EFCE", "complete": "C6EFCE",
    "running": "FFEB9C", "submitted": "FFEB9C", "partial": "FFEB9C", "report incomplete": "FFEB9C",
    "failed": "FFC7CE", "not run": "F2F2F2", "no": "F2F2F2", "not started": "F2F2F2",
}


def build(out, rows, root):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "tracker"
    gen = wb.create_sheet("_generated")
    gen.sheet_state = "hidden"
    headers = [h for h, _, _ in COLUMNS]
    ws.append(headers)
    gen.append(headers)
    for row in rows:
        vals = [row.get(k, "") for _, k, _ in COLUMNS]
        ws.append(vals)
        gen.append(vals)
    bold = Font(bold=True, color="FFFFFF")
    for c in ws[1]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="44546A")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, (_, key, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        if key in ("farm_status", "data_local", "review_progress", "review_signed_off"):
            for r in range(2, ws.max_row + 1):
                c = ws.cell(r, i)
                f = FILL.get(str(c.value))
                if f:
                    c.fill = PatternFill("solid", fgColor=f)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in ("E", "T"))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    col = {k: get_column_letter(i) for i, (_, k, _) in enumerate(COLUMNS, 1)}
    last = ws.max_row
    dv = DataValidation(type="list", formula1='"' + ",".join(FARM_STATUS) + '"', allow_blank=True)
    dv.add(f"{col['farm_status']}2:{col['farm_status']}{last}")
    yn = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    yn.add(f"{col['review_signed_off']}2:{col['review_signed_off']}{last}")
    ws.add_data_validation(dv)
    ws.add_data_validation(yn)

    s = wb.create_sheet("summary", 1)
    s.append(["Run tracker", datetime.datetime.now().strftime("refreshed %Y-%m-%d %H:%M")])
    s["A1"].font = Font(bold=True, size=13)
    s.append([])
    s.append(["", "patients", "farm run done", "data local", "review signed off"])
    for c in s[3]:
        c.font = Font(bold=True)
    groups = [("TP53 clade", lambda r: "TP53" in r["reason"]),
              ("BCR::ABL1", lambda r: "BCR::ABL1" in r["reason"]),
              ("transplant pair", lambda r: "transplant" in r["reason"]),
              ("all", lambda r: True)]
    for name, pred in groups:
        g = [r for r in rows if pred(r)]
        s.append([name, len(g), sum(r["farm_status"] == "done" for r in g),
                  sum(r["data_local"] == "yes" for r in g), sum(r["review_signed_off"] == "yes" for r in g)])
    s.append([])
    s.append(["Manual columns (edit here, they survive a rebuild): Farm run, Farm run dir, PEAR-TREE commit, "
              "Results dir, Review signed off, Notes."])
    s.append(["Automatic columns (rebuilt from ~/Documents/hsc_results): Colonies, Data local, Somatic rows, "
              "Tiers, Excluded, Reviewed (insertion_review verdicts), Review progress."])
    s.append(["A commit starting with '~' was inferred from dates, not recorded by the run; runs submitted after "
              "pipeline.sh started stamping PT_COMMIT into run.env show the real one."])
    s.append([f"Rebuild: python3 {os.path.relpath(__file__, REPO)}   (roster: patients/run_tracker.tsv)"])
    s.column_dimensions["A"].width = 18
    for c in "BCDE":
        s.column_dimensions[c].width = 16
    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tsv", default=os.path.join(REPO, "patients", "run_tracker.tsv"))
    ap.add_argument("--out", default=os.path.join(REPO, "patients", "run_tracker.xlsx"))
    ap.add_argument("--results", default=os.path.expanduser("~/Documents/hsc_results"))
    a = ap.parse_args()

    comments, rows = read_tsv(a.tsv)
    edits = merge_back(a.out, rows)
    for pd, key, old, new in edits:
        print(f"xlsx edit kept: {pd} {key}: {old!r} -> {new!r}")
    if edits:
        write_tsv(a.tsv, comments, rows)
    full = [dict(r, **auto_fields(a.results, r)) for r in rows]
    build(a.out, full, a.results)
    print(f"wrote {a.out} ({len(full)} patients)")


if __name__ == "__main__":
    sys.exit(main())
